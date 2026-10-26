"""Mẫu FIS đợt 5 (2026-10-26): nhập đội + VĐV cấp giải từ Excel, xuất kết quả Excel."""
import io
import openpyxl
from tests.test_batch19_team_format import _create_team_tournament, _h


def _xlsx(rows):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Đội", "Họ và tên", "Giới tính", "Hạng", "Mã TV", "SĐT"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


def test_template_and_import_creates_teams_members_guests(client, world):
    h = _h(world)
    r = client.get("/api/team-import/template", headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    t = _create_team_tournament(client, world)
    tid = t["id"]
    m0, m1, m2 = world.members1[0], world.members1[1], world.members1[2]
    content = _xlsx([
        ["Đội Rồng", m0.full_name, "Nam", "A", m0.member_code, ""],       # khớp mã TV, ghi đè hạng A
        ["Đội Rồng", m1.full_name, "", "", "", ""],                        # khớp tên duy nhất
        ["Đội Hổ", "Khách Mời Mới", "Nữ", "C", "", "0999000111"],          # không khớp → tạo khách
        ["", m2.full_name, "", "", "", ""],                                 # không đội → chờ phân đội
        ["Đội Hổ", "", "", "", "", ""],                                     # thiếu tên → lỗi
        ["Đội Hổ", "Ai Đó", "", "", "TV9999", ""],                          # mã không tồn tại → lỗi
        ["Đội Rồng", m0.full_name, "", "", m0.member_code, ""],            # trùng → skipped
    ])
    r = client.post(f"/api/tournaments/{tid}/team-players/import", files={"file": ("vdv.xlsx", content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=h)
    assert r.status_code == 200, r.text
    j = r.json()
    assert sorted(j["teams_created"]) == ["Đội Hổ", "Đội Rồng"]
    assert j["added"] == 4 and j["guests_created"] == 1 and j["pool_count"] == 1
    assert len(j["errors"]) == 2 and j["skipped"] == []   # dòng trùng trong cùng file được gộp, không báo bỏ qua
    teams = {x["participant"]["team_name"]: x for x in client.get(f"/api/tournaments/{tid}/teams", headers=h).json()}
    rong = teams["Đội Rồng"]["players"]
    assert len(rong) == 2 and next(p for p in rong if p["member_id"] == m0.id)["rank"] == "A"
    ho = teams["Đội Hổ"]["players"]
    assert len(ho) == 1 and ho[0]["player_id"] and ho[0]["gender"] == "female" and ho[0]["rank"] == "C"
    # Nhập lại cùng file: không tạo thêm đội/khách, mọi người bị bỏ qua
    r2 = client.post(f"/api/tournaments/{tid}/team-players/import", files={"file": ("vdv.xlsx", content,
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=h).json()
    assert r2["teams_created"] == [] and r2["guests_created"] == 0 and r2["added"] == 0 and len(r2["skipped"]) == 4
    # Quyền xem không nhập được
    from tests.conftest import World
    assert client.post(f"/api/tournaments/{tid}/team-players/import", files={"file": ("v.xlsx", content, "application/octet-stream")},
                       headers=World.headers(world.viewer1, world.club1)).status_code == 403


def test_export_tournament_excel_sheets(client, world, db):
    from tests.test_batch19_team_format import _two_teams
    t = _create_team_tournament(client, world)
    tid = t["id"]; h = _h(world)
    _two_teams(client, world, db, tid)
    client.put(f"/api/tournaments/{tid}", json={"status": "active"}, headers=h)
    client.post(f"/api/tournaments/{tid}/generate", headers=h)
    r = client.get(f"/api/tournaments/{tid}/export.xlsx", headers=h)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tổng quan", "Đội & VĐV", "Trận đấu", "Xếp hạng"]
    assert wb["Đội & VĐV"].max_row >= 9          # 8 VĐV + header
    assert wb["Trận đấu"].max_row >= 2
    # Giải đơn: không có sheet Đội & VĐV
    t2 = client.post("/api/tournaments", json={"name": "Đơn", "format": "round_robin", "team_type": "singles",
                     "member_ids": [m.id for m in world.members1[:3]]}, headers=h).json()
    r2 = client.get(f"/api/tournaments/{t2['id']}/export.xlsx", headers=h)
    assert r2.status_code == 200 and "Đội & VĐV" not in openpyxl.load_workbook(io.BytesIO(r2.content)).sheetnames
